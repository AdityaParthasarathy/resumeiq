"""One-off script to (re)generate static test fixture files.

Not part of the app or its runtime dependencies -- run manually if the
fixtures under tests/fixtures/ ever need to change.
"""

import os

from docx import Document
from fpdf import FPDF

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures")

SAMPLE_TEXT_LINES = [
    "Jordan Lee",
    "jordan.lee@example.com | (555) 123-4567",
    "",
    "Skills",
    "Python, SQL, Excel, JavaScript, Communication",
    "",
    "Education",
    "B.S. in Computer Science, State University, 2023",
    "",
    "Projects",
    "Built a sales dashboard using Python and SQL to track quarterly revenue.",
]

STRONG_TEXT_LINES = [
    "Alex Morgan",
    "alex.morgan@example.com | (555) 987-6543 | linkedin.com/in/alexmorgan | github.com/alexmorgan",
    "",
    "Summary",
    "Results-driven software engineer with 3 years of experience building scalable web applications.",
    "",
    "Skills",
    "Python, JavaScript, React, SQL, AWS, Docker, Git, REST APIs, Machine Learning, Agile",
    "",
    "Education",
    "B.S. in Computer Science, University of Springfield, 2021",
    "",
    "Experience",
    "Software Engineer, Acme Corp, 2021-2024",
    "- Led a team of 4 engineers to redesign the checkout flow, increasing conversion by 18%",
    "- Built and deployed a microservices architecture using Docker and AWS, reducing latency by 35%",
    "- Optimized database queries, cutting average response time from 800ms to 200ms",
    "- Responsible for maintaining CI/CD pipelines using GitHub Actions",
    "",
    "Projects",
    "Personal Finance Tracker",
    "- Developed a full-stack budgeting app using React and Flask, used by over 500 users",
    "- Implemented JWT-based authentication and role-based access control",
    "Open Source Contribution",
    "- Contributed a caching layer to an open-source Python library, reducing memory usage by 20%",
]


def generate_docx():
    document = Document()
    for line in SAMPLE_TEXT_LINES:
        document.add_paragraph(line)
    document.save(os.path.join(FIXTURES_DIR, "sample_resume.docx"))


def generate_strong_resume_docx():
    document = Document()
    for line in STRONG_TEXT_LINES:
        document.add_paragraph(line)
    document.save(os.path.join(FIXTURES_DIR, "strong_resume.docx"))


def generate_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in SAMPLE_TEXT_LINES:
        pdf.cell(0, 8, text=line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(os.path.join(FIXTURES_DIR, "sample_resume.pdf"))


def generate_empty_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.output(os.path.join(FIXTURES_DIR, "empty.pdf"))


if __name__ == "__main__":
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    generate_docx()
    generate_strong_resume_docx()
    generate_pdf()
    generate_empty_pdf()
    print("Fixtures written to", os.path.abspath(FIXTURES_DIR))
